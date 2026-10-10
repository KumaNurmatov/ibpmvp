"""Проверки книги: раскладка, формулы и их посчитанные значения."""
from __future__ import annotations

from openpyxl import load_workbook

from tool.pipeline import export_xlsx, ledger
from tool.tests import fixtures as fx


def build(extractions):
    orders, orphans = ledger.build(list(extractions))
    return orders, orphans


def write(tmp_path, extractions, name="книга.xlsx"):
    orders, orphans = build(extractions)
    path = export_xlsx.write(orders, tmp_path / name, orphans=orphans)
    return path, orders


def test_листы_как_в_ручной_книге(tmp_path):
    path, _ = write(tmp_path, fx.ORDER_666)
    wb = load_workbook(path)
    assert wb.sheetnames == ["недоплата", "нет постоплат", "закрытые", "проверки"]


def test_незакрытый_заказ_без_постоплаты_попадает_на_свой_лист(tmp_path):
    path, _ = write(tmp_path, fx.ORDER_666)
    wb = load_workbook(path)
    assert wb["нет постоплат"]["A4"].value == "Заказ Аи 13.04.26/4"
    assert wb["нет постоплат"]["D4"].value == "666/1-1"
    assert wb["недоплата"]["A4"].value is None


def test_оплаченный_заказ_уезжает_в_закрытые(tmp_path):
    path, _ = write(tmp_path, list(fx.ORDER_666) + [fx.POSTPAY])
    wb = load_workbook(path)
    assert wb["закрытые"]["A4"].value == "Заказ Аи 13.04.26/4"


def test_формулы_на_тех_же_местах(tmp_path):
    path, _ = write(tmp_path, fx.ORDER_666)
    ws = load_workbook(path)["нет постоплат"]
    assert ws["B8"].value == "=ROUND(B6*B7,2)"
    assert ws["E12"].value == "=ROUND(B7*B12,2)"
    assert ws["J11"].value.startswith("=SUM(J6:J10)-SUMIF(H6:H10,")
    assert ws["J12"].value.startswith("=SUMIF(H6:H10,")
    assert ws["J14"].value == '=IF(E12="","",ROUND(E12-J11-J12,2))'


def test_у_формул_есть_посчитанное_значение(tmp_path):
    """Без кэша ячейка выглядит пустой везде, где файл не пересчитывается."""
    path, _ = write(tmp_path, fx.ORDER_666)
    ws = load_workbook(path, data_only=True)["нет постоплат"]
    assert ws["B8"].value == 300000.0
    assert ws["E12"].value == 324800.0
    assert ws["J11"].value == 150000.0
    assert ws["J14"].value == 174800.0


def test_в_карточке_отгрузка_и_платежи(tmp_path):
    path, _ = write(tmp_path, fx.ORDER_666)
    ws = load_workbook(path)["нет постоплат"]
    assert ws["B6"].value == 750          # заказано
    assert ws["B7"].value == 400          # цена
    assert ws["B12"].value == 812         # отгружено
    assert ws["C12"].value == 8           # мест
    assert ws["D12"].value == 282         # вес
    assert ws["H6"].value == "Предоплата"
    assert ws["J6"].value == 150000.0


def test_замечания_собраны_на_отдельном_листе(tmp_path):
    path, orders = write(tmp_path, fx.ORDER_666)
    ws = load_workbook(path)["проверки"]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    codes = {r[3] for r in rows if r[3]}
    assert {"overdelivery", "late_shipment", "conflict"} <= codes
    assert all(r[1] == "666/1-1" for r in rows if r[1])


def test_блоки_не_наезжают_друг_на_друга(tmp_path):
    """Второй заказ должен начинаться ровно через 15 строк."""
    path, _ = write(tmp_path, [fx.SPECIFICATION, fx.SHIPMENT, fx.STATEMENT])
    wb = load_workbook(path)
    ws = wb["недоплата"]
    starts = [c.row for c in ws["A"] if str(c.value or "").startswith("Заказ ")]
    assert len(starts) >= 2
    assert all(b - a == export_xlsx.BLOCK_ROWS for a, b in zip(starts, starts[1:]))


def test_неприкаянные_файлы_видны_в_проверках(tmp_path):
    stray = dict(fx.SHIPMENT, contract_no=None, order_no=None)
    orders, orphans = build([stray])
    path = export_xlsx.write(orders, tmp_path / "x.xlsx", orphans=orphans)
    ws = load_workbook(path)["проверки"]
    text = " ".join(str(c.value) for row in ws.iter_rows(min_row=2) for c in row)
    assert "отправка 666.pdf" in text and "unrouted" in text
