"""Проверки первого слоя: файл → страницы."""
from __future__ import annotations

import docx
import pytest
from openpyxl import Workbook
from PIL import Image

from tool.pipeline import docs


def test_текстовый_файл(tmp_path):
    path = tmp_path / "заметка.txt"
    path.write_text("предоплата 150000", encoding="utf-8")
    doc = docs.load(path)
    assert doc.kind == "text"
    assert "150000" in doc.text


def test_docx_забирает_и_таблицы(tmp_path):
    path = tmp_path / "реестр.docx"
    d = docx.Document()
    d.add_paragraph("Реестр постоплат")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Договор"
    table.cell(0, 1).text = "666/1-1"
    table.cell(1, 0).text = "Остаток"
    table.cell(1, 1).text = "174800"
    d.save(path)
    doc = docs.load(path)
    assert doc.kind == "docx"
    assert "666/1-1" in doc.text and "174800" in doc.text


def test_xlsx_постранично_по_листам(tmp_path):
    path = tmp_path / "книга.xlsx"
    wb = Workbook()
    wb.active.title = "недоплата"
    wb.active["A1"] = "Заказ Аи 13.04.26/4"
    wb.create_sheet("нет постоплат")["A1"] = "Заказ Аи 01.04.26/1"
    wb.save(path)
    doc = docs.load(path)
    assert doc.kind == "xlsx"
    assert len(doc.pages) == 2
    assert "недоплата" in doc.pages[0].text


def test_картинка_едет_в_base64(tmp_path):
    path = tmp_path / "накладная.png"
    Image.new("RGB", (40, 30), "white").save(path)
    doc = docs.load(path)
    assert doc.kind == "image"
    assert doc.has_images and doc.pages[0].image_b64


def test_неподдерживаемый_формат_не_роняет_разбор(tmp_path):
    path = tmp_path / "архив.zip"
    path.write_bytes(b"PK\x03\x04")
    doc = docs.load(path)
    assert doc.kind == "unsupported"
    assert doc.error


def test_битый_файл_не_роняет_разбор(tmp_path):
    path = tmp_path / "сломано.docx"
    path.write_bytes("не docx вовсе".encode())
    doc = docs.load(path)
    assert doc.kind == "unsupported"
    assert doc.error


@pytest.mark.parametrize("name", ["скан.pdf", "СКАН.PDF"])
def test_регистр_расширения_не_важен(tmp_path, name):
    path = tmp_path / name
    path.write_bytes("%PDF-1.4\n%битый".encode())
    doc = docs.load(path)
    assert doc.kind in {"pdf_text", "pdf_scan", "unsupported"}
