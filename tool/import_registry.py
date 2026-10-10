#!/usr/bin/env python3
"""Перенос реестра Excel в data/orders.json.

Реестр — книга с листами «недоплата» и «нет постоплат», где каждый заказ
занимает ровно 15 строк: номер и договор, изделие, заказано и цена, оплаты
справа, отправка снизу, в блоке «Учет» — названия файлов со ссылками на
Google Drive.

Переносится только то, что в книге написано. Считаемые ячейки (сумма заказа,
оплачено, остаток) не переносятся: их считает сайт, и расхождение с книгой —
это находка, а не повод переписать цифру. Скрипт сам сверяет свой расчёт с
остатком из книги и показывает все расхождения.

    python3 tool/import_registry.py реестр.xlsx web/data/orders.json
    python3 tool/import_registry.py реестр.xlsx web/data/orders.json --dry-run
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

BLOCK = 15                      # высота блока одного заказа
PAY_OFFSETS = (2, 3, 4, 5)      # строки оплат: H вид, I дата, J сумма
LINK_OFFSETS = (10, 11, 12, 13) # строки «Учет»: A подпись, B имя файла со ссылкой
SHIP = 8                        # строка отправки: A дата, B кол-во, C мест, D вес
TRANSFERS = (8, 10)             # «Подтвержденные переносы», столбец J
REMAINDER = (10, 10)            # «РАСЧЁТНЫЙ ОСТАТОК» из книги — только для сверки

KINDS = {
    "предоплата": "предоплата",
    "постоплата": "постоплата",
    "доп. постоплата": "доп. постоплата",
    "доп постоплата": "доп. постоплата",
    "перенос": "перенос",
}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("registry", type=Path, help="книга Excel с реестром")
    ap.add_argument("target", type=Path, help="файл data/orders.json")
    ap.add_argument("--dry-run", action="store_true", help="только показать, ничего не писать")
    args = ap.parse_args(argv)

    orders, problems = read_registry(args.registry)
    print(f"Прочитано заказов: {len(orders)}")

    existing = json.loads(args.target.read_text("utf-8")) if args.target.exists() else {"orders": []}
    merged, kept, replaced = merge(existing.get("orders", []), orders)
    print(f"Было: {len(existing.get('orders', []))} · добавлено: {len(merged) - kept} · "
          f"обновлено: {replaced} · стало: {len(merged)}")

    for line in problems:
        print("  ·", line)

    if args.dry_run:
        print("Пробный прогон: файл не тронут.")
        return 0

    existing["orders"] = merged
    existing.setdefault("version", 1)
    existing["updated"] = dt.date.today().isoformat()
    args.target.write_text(json.dumps(existing, ensure_ascii=False, indent=2) + "\n", "utf-8")
    print(f"Записано в {args.target}")
    return 0


def read_registry(path: Path) -> tuple[list[dict], list[str]]:
    book = openpyxl.load_workbook(path, data_only=True)
    orders: list[dict] = []
    problems: list[str] = []

    for sheet in book:
        for top in block_starts(sheet):
            order, note = read_block(sheet, top)
            if order:
                orders.append(order)
            if note:
                problems.append(note)
    return orders, problems


def block_starts(sheet) -> list[int]:
    return [r for r in range(1, sheet.max_row + 1)
            if str(sheet.cell(r, 1).value or "").startswith("Заказ ")]


def read_block(sheet, top: int) -> tuple[dict | None, str | None]:
    cell = lambda off, col: sheet.cell(top + off, col).value  # noqa: E731

    contract = text(cell(0, 4))
    order_no = re.sub(r"^Заказ\s+", "", text(cell(0, 1))).strip()
    key = digits(contract) or digits(order_no) or order_no
    if not key:
        return None, f"{sheet.title}, строка {top}: нет ни договора, ни номера заказа — пропущено"

    price = number(cell(3, 2))
    order = {
        "id": key,
        "order_no": order_no,
        "contract_no": contract,
        # Дата заказа — та, что стоит в его же номере: «Аи 29.01.26/2».
        "order_date": date_from_number(order_no),
        "item": text(cell(1, 1)),
        "ordered_qty": number(cell(2, 2)),
        "unit_price": price,
        "tnved": "",
        "item_article": "",
        "workshops": [],
        "photo": "",
        "shipment": {
            "date": date(cell(SHIP, 1)),
            "qty": number(cell(SHIP, 2)),
            "places": number(cell(SHIP, 3)),
            "weight_gross": number(cell(SHIP, 4)),
            # Водителя и его телефон реестр не хранит — они только в CMR.
            "driver": "",
            "driver_phone": "",
        },
        # Приёмка в книге не велась: заполняется руками на сайте.
        "acceptance": {"accepted_qty": None, "rejected_qty": None, "date": "", "reason": ""},
        "payments": payments(sheet, top),
        # В книге условий нет: проценты предоплаты и сроки живут в договоре.
        # Ставим то, что известно (половина вперёд), остальное — пусто, иначе
        # сайт посчитал бы просрочку и неустойку по выдуманному сроку.
        "terms": {"prepay_percent": 50, "lead_time_days": None, "postpay_days": None},
        "links": links(sheet, top),
        "note": "",
    }

    return order, mismatch(order, sheet, top, price)


def payments(sheet, top: int) -> list[dict]:
    out = []
    for off in PAY_OFFSETS:
        kind = text(sheet.cell(top + off, 8).value).lower()
        amount = number(sheet.cell(top + off, 10).value)
        if not amount:                       # пустая или нулевая строка — не платёж
            continue
        out.append({
            "kind": KINDS.get(kind, "неизвестно"),
            "date": date(sheet.cell(top + off, 9).value),
            "amount": amount,
            "doc_number": "",
            "purpose": "",
        })

    transfers = number(sheet.cell(top + TRANSFERS[0], TRANSFERS[1]).value)
    if transfers:
        out.append({"kind": "перенос", "date": "", "amount": transfers,
                    "doc_number": "", "purpose": "Подтверждённый перенос из реестра"})
    return out


def links(sheet, top: int) -> list[dict]:
    out = []
    for off in LINK_OFFSETS:
        cell = sheet.cell(top + off, 2)
        url = cell.hyperlink.target if cell.hyperlink else ""
        if not url:
            continue
        title = text(cell.value) or text(sheet.cell(top + off, 1).value) or "Документ"
        out.append({"title": title, "url": url})
    return out


def mismatch(order: dict, sheet, top: int, price: float | None) -> str | None:
    """Сверка с остатком, посчитанным в книге: расхождение — это находка."""
    stated = number(sheet.cell(top + REMAINDER[0], REMAINDER[1]).value)
    qty = order["shipment"]["qty"]
    if stated is None or qty is None or price is None:
        return None
    paid = sum(p["amount"] for p in order["payments"])
    ours = round(qty * price - paid, 2)
    if abs(ours - stated) <= 0.5:
        return None
    return (f"{order['order_no']} (договор {order['contract_no']}): в книге остаток "
            f"{stated:,.2f}, по её же цифрам выходит {ours:,.2f} — "
            f"разница {ours - stated:,.2f}".replace(",", " "))


def merge(old: list[dict], new: list[dict]) -> tuple[list[dict], int, int]:
    """Карточки с тем же ключом заменяются, остальные остаются на месте."""
    by_key = {o["id"]: o for o in old}
    kept = len(by_key)
    replaced = 0
    for order in new:
        if order["id"] in by_key:
            replaced += 1
        by_key[order["id"]] = order
    return list(by_key.values()), kept, replaced


# ── мелочи ────────────────────────────────────────────────────────────────
def date_from_number(order_no: str) -> str:
    """«Аи 29.01.26/2» → 2026-01-29. Год бывает и четырёхзначным."""
    m = re.search(r"(\d{2})\.(\d{2})\.(\d{4}|\d{2})(?!\d)", order_no or "")
    if not m:
        return ""
    day, month, year = (int(x) for x in m.groups())
    try:
        return dt.date(year if year > 100 else 2000 + year, month, day).isoformat()
    except ValueError:
        return ""


def text(value) -> str:
    return "" if value is None else str(value).strip()


def digits(value: str) -> str:
    match = re.search(r"\d{2,6}", value or "")
    return match.group(0) if match else ""


def number(value) -> float | None:
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    if isinstance(value, str):
        cleaned = value.replace("\xa0", "").replace(" ", "").replace(",", ".")
        try:
            return round(float(cleaned), 2)
        except ValueError:
            return None
    return None


def date(value) -> str:
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return ""


if __name__ == "__main__":
    sys.exit(main())
