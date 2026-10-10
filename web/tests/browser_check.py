"""Сквозная проверка браузерной версии в настоящем Chromium.

Проверяет то, чего не видят модульные тесты: читает ли pdf.js настоящие PDF
(текстовый слой и скан), собирает ли ExcelJS книгу с формулами и их
значениями. Книга потом открывается openpyxl и сверяется по ячейкам.

    python3 web/tests/browser_check.py <папка-с-pdf>
"""
from __future__ import annotations

import base64
import json
import shutil
import subprocess
import sys
import tempfile
import threading
import time
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
REPO = Path(__file__).resolve().parents[2]
PORT = 8123


def serve(directory: Path) -> ThreadingHTTPServer:
    handler = partial(SimpleHTTPRequestHandler, directory=str(directory))
    server = ThreadingHTTPServer(("127.0.0.1", PORT), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server


def main(pdf_dir: Path) -> int:
    from playwright.sync_api import sync_playwright

    stage = Path(tempfile.mkdtemp())
    shutil.copytree(REPO / "web", stage / "web")
    (stage / "pdf").mkdir()
    for pdf in sorted(pdf_dir.glob("*.pdf")):
        shutil.copy(pdf, stage / "pdf" / pdf.name)

    server = serve(stage)
    failures: list[str] = []
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(executable_path=CHROME)
            page = browser.new_page()
            page.on("console", lambda m: print("  [консоль]", m.type, m.text)
                    if m.type == "error" else None)
            page.on("pageerror", lambda e: failures.append(f"ошибка страницы: {e}"))
            page.goto(f"http://127.0.0.1:{PORT}/web/index.html", wait_until="networkidle")

            print("— pdf.js на настоящих файлах")
            for name in sorted(x.name for x in (stage / "pdf").iterdir()):
                info = page.evaluate(ASSESS_PDF, f"/pdf/{name}")
                size_kb = round(info["b64"] / 1024)
                print(f"  {name:18} {info['kind']:9} страниц {info['pages']:2} "
                      f"картинок {info['images']:2} base64 {size_kb:5} КБ "
                      f"текста {info['chars']:6}")
                if info["kind"] == "unsupported":
                    failures.append(f"{name}: pdf.js не справился — {info.get('error')}")

            print("— сведение и книга на фикстурах")
            result = page.evaluate(RUN_FIXTURES)
            print(f"  заказов {result['orders']}, остаток {result['remainder']}, "
                  f"замечаний {result['flags']}")
            if result["orders"] != 1 or result["remainder"] != 174800:
                failures.append(f"расчёт разошёлся: {result}")

            book = base64.b64decode(page.evaluate(BUILD_BOOK))
            out = Path("/tmp/браузерная-книга.xlsx")
            out.write_bytes(book)
            print(f"  книга собрана: {len(book)} байт → {out}")
            browser.close()
    finally:
        server.shutdown()
        shutil.rmtree(stage, ignore_errors=True)

    failures += verify_workbook(out)
    print()
    if failures:
        print("ПРОВАЛЫ:")
        for f in failures:
            print("  ✗", f)
        return 1
    print("всё сошлось")
    return 0


def verify_workbook(path: Path) -> list[str]:
    """Книга из браузера должна быть неотличима от серверной по структуре."""
    from openpyxl import load_workbook

    bad: list[str] = []
    wb = load_workbook(path)
    wv = load_workbook(path, data_only=True)
    print("— книга из браузера")
    print("  листы:", wb.sheetnames)
    if wb.sheetnames != ["недоплата", "нет постоплат", "закрытые", "проверки"]:
        bad.append(f"листы не те: {wb.sheetnames}")

    ws, wsv = wb["нет постоплат"], wv["нет постоплат"]
    expect_formula = {
        "B8": "ROUND(B6*B7,2)",
        "E12": "ROUND(B7*B12,2)",
        "J14": 'IF(E12="","",ROUND(E12-J11-J12,2))',
    }
    for ref, formula in expect_formula.items():
        got = str(ws[ref].value or "").lstrip("=")
        if got != formula:
            bad.append(f"{ref}: формула {got!r}, ожидали {formula!r}")
    expect_value = {"B6": 750, "B7": 400, "B8": 300000, "B12": 812, "C12": 8,
                    "D12": 282, "E12": 324800, "J11": 150000, "J14": 174800}
    for ref, value in expect_value.items():
        got = wsv[ref].value
        print(f"  {ref:4} {str(ws[ref].value)[:34]:36} → {got}")
        if got != value:
            bad.append(f"{ref}: значение {got!r}, ожидали {value!r}")

    notes = sum(1 for row in ws.iter_rows() for c in row if c.comment)
    print("  примечаний с источниками:", notes)
    if notes < 4:
        bad.append(f"примечаний всего {notes} — источники потерялись")

    codes = {r[3].value for r in wb["проверки"].iter_rows(min_row=2) if r[3].value}
    print("  коды проверок:", sorted(codes))
    if not {"overdelivery", "late_shipment", "conflict"} <= codes:
        bad.append(f"на листе проверок не хватает находок: {sorted(codes)}")
    return bad


ASSESS_PDF = """
async (url) => {
  const { load } = await import("/web/lib/docs.js");
  const blob = await (await fetch(url)).blob();
  const doc = await load(new File([blob], url.split("/").pop(), { type: "application/pdf" }));
  return {
    kind: doc.kind, pages: doc.pages.length, error: doc.error,
    images: doc.pages.filter((p) => p.image_b64).length,
    b64: doc.pages.reduce((s, p) => s + (p.image_b64 ? p.image_b64.length : 0), 0),
    chars: doc.pages.reduce((s, p) => s + (p.text ? p.text.length : 0), 0),
  };
}
"""

# Фикстуры грузим модулем, подменяя распознавание: проверяем сведение и книгу,
# а не модель.
RUN_FIXTURES = """
async () => {
  const fx = await import("/web/tests/fixtures.js");
  const stub = { extract: async (doc) => fx.ORDER_666.find((e) => e._source_file === doc.filename) };
  const files = fx.ORDER_666.map((e) =>
    new File([new Uint8Array([1, 2, 3])], e._source_file, { type: "application/pdf" }));
  // load() на таком файле вернёт unsupported, поэтому зовём сведение напрямую.
  const { orders, orphans } = window.__ibp.buildOrders(fx.ORDER_666);
  window.__orders = orders;
  window.__orphans = orphans;
  return {
    orders: orders.length, orphans: orphans.length,
    remainder: orders[0].computed.remainder,
    flags: orders[0].flags.length,
  };
}
"""

BUILD_BOOK = """
async () => {
  const blob = await window.__ibp.buildWorkbook(window.__orders, window.__orphans);
  const buf = new Uint8Array(await blob.arrayBuffer());
  let s = "";
  for (let i = 0; i < buf.length; i += 0x8000) {
    s += String.fromCharCode.apply(null, buf.subarray(i, i + 0x8000));
  }
  return btoa(s);
}
"""


if __name__ == "__main__":
    directory = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    if not directory or not directory.is_dir():
        print(__doc__)
        raise SystemExit(2)
    raise SystemExit(main(directory))
